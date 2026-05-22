import csv
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Set, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


class StaffCompareError(ValueError):
    pass


def _read_rows(path: Path) -> List[List[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "gbk"):
            try:
                with open(path, "r", encoding=encoding, newline="") as f:
                    return [[str(cell or "").strip() for cell in row] for row in csv.reader(f)]
            except UnicodeDecodeError:
                continue
        raise StaffCompareError("CSV 文件编码无法识别")

    if suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            return [[str(cell.value or "").strip() for cell in row] for row in ws.iter_rows()]
        finally:
            wb.close()

    raise StaffCompareError("仅支持 .xlsx / .csv 文件")


def _looks_like_header(row: List[str], keywords: Tuple[str, ...]) -> bool:
    joined = "".join(row[:3])
    return any(keyword in joined for keyword in keywords)


def _active_staff(path: Path) -> Tuple[Set[str], Dict[str, str]]:
    rows = _read_rows(path)
    names: Set[str] = set()
    departments: Dict[str, str] = {}
    for index, row in enumerate(rows):
        if index == 0 and _looks_like_header(row, ("姓名", "人员", "部门")):
            continue
        name = row[0].strip() if len(row) > 0 else ""
        department = row[1].strip() if len(row) > 1 else ""
        if name:
            names.add(name)
            departments[name] = department
    if not names:
        raise StaffCompareError("在职人员表未识别到姓名")
    return names, departments


def _account_rows(path: Path) -> List[Dict[str, str]]:
    rows = _read_rows(path)
    accounts: List[Dict[str, str]] = []
    for index, row in enumerate(rows):
        if index == 0 and _looks_like_header(row, ("姓名", "邮箱", "账号")):
            continue
        name = row[0].strip() if len(row) > 0 else ""
        email = row[1].strip() if len(row) > 1 else ""
        if name or email:
            accounts.append({"姓名": name, "邮箱账号": email})
    if not accounts:
        raise StaffCompareError("邮箱账号表未识别到账号数据")
    return accounts


def _style(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 36)

    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName="RetiredEmailAccounts", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def create_staff_compare_workbook(active_path: Path, accounts_path: Path, output_dir: Path) -> Dict:
    active_names, active_departments = _active_staff(active_path)
    accounts = _account_rows(accounts_path)
    suspicious = [row for row in accounts if row["姓名"] not in active_names]

    wb = Workbook()
    ws = wb.active
    ws.title = "疑似离职退休邮箱"
    ws.append(["姓名", "邮箱账号", "对比结果", "在职部门"])
    for row in suspicious:
        ws.append([row["姓名"], row["邮箱账号"], "邮箱表存在但在职表不存在", active_departments.get(row["姓名"], "")])

    _style(ws)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"疑似离职退休人员_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(output_path)

    return {
        "output_path": output_path,
        "active_count": len(active_names),
        "account_count": len(accounts),
        "suspicious_count": len(suspicious),
    }
