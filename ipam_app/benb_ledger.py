import csv
import ipaddress
import re
from pathlib import Path
from typing import Dict, List, Set

from openpyxl import load_workbook

import config


LEDGER_FILE = config.DATA_DIR / "BenB.csv"
POOL_FILE = config.DATA_DIR / "BenB_pool.csv"
LEDGER_COLUMNS = ["IP地址", "状态", "姓名", "账号"] + [f"预留{i}" for i in range(1, 14)] + ["房间号"]
IP_PATTERN = re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b")


class BenBLedgerError(ValueError):
    pass


def _read_csv_rows(path: Path) -> List[List[str]]:
    if not path.exists():
        return []
    for encoding in ("utf-8-sig", "gbk"):
        try:
            with open(path, "r", encoding=encoding, newline="") as f:
                return [[str(cell or "").strip() for cell in row] for row in csv.reader(f)]
        except UnicodeDecodeError:
            continue
    raise BenBLedgerError("BenB.csv 编码无法识别")


def _normalize_row(row: List[str]) -> Dict:
    values = list(row) + [""] * max(0, 18 - len(row))
    room = values[17] or (values[4] if len(row) <= 5 else "")
    return {
        "IP地址": values[0].strip(),
        "状态": values[1].strip(),
        "姓名": values[2].strip(),
        "账号": values[3].strip(),
        "房间号": room.strip(),
    }


def read_ledger() -> List[Dict]:
    rows = _read_csv_rows(LEDGER_FILE)
    if not rows:
        return []
    data_rows = rows[1:] if rows and _looks_like_header(rows[0]) else rows
    return [_normalize_row(row) for row in data_rows if any(str(cell).strip() for cell in row)]


def _looks_like_header(row: List[str]) -> bool:
    joined = "".join(row[:5])
    return "IP" in joined and ("姓名" in joined or "账号" in joined or "状态" in joined)


def _row_to_csv(row: Dict) -> List[str]:
    values = [""] * 18
    values[0] = row.get("IP地址", "").strip()
    values[1] = row.get("状态", "").strip()
    values[2] = row.get("姓名", "").strip()
    values[3] = row.get("账号", "").strip()
    values[17] = row.get("房间号", "").strip()
    return values


def write_ledger(rows: List[Dict]) -> None:
    config.DATA_DIR.mkdir(parents=True, exist_ok=True)
    with open(LEDGER_FILE, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(LEDGER_COLUMNS)
        for row in sorted(rows, key=lambda item: _ip_sort_key(item.get("IP地址", ""))):
            writer.writerow(_row_to_csv(row))


def _ip_sort_key(value: str) -> int:
    try:
        return int(ipaddress.IPv4Address(value.strip()))
    except ValueError:
        return 2**32


def _validate_ip(ip_value: str) -> str:
    try:
        return str(ipaddress.IPv4Address(ip_value.strip()))
    except ValueError:
        raise BenBLedgerError(f"无效 IP 地址: {ip_value}")


def search(keyword: str) -> List[Dict]:
    keyword = (keyword or "").strip()
    if not keyword:
        return []
    rows = read_ledger()
    lowered = keyword.lower()
    results = []
    for row in rows:
        if (
            lowered in row.get("IP地址", "").lower()
            or lowered in row.get("姓名", "").lower()
            or lowered in row.get("账号", "").lower()
            or lowered in row.get("房间号", "").lower()
            or lowered in row.get("状态", "").lower()
        ):
            results.append(row)
    return results


def batch_search(text: str) -> List[Dict]:
    keywords = [part.strip() for part in re.split(r"[\s,;；，]+", text or "") if part.strip()]
    output = []
    seen = set()
    for keyword in keywords:
        matches = search(keyword)
        if not matches:
            output.append({"查询关键字": keyword, "IP地址": "", "状态": "", "姓名": "", "账号": "", "房间号": "", "结果": "未找到"})
            continue
        for row in matches:
            key = (keyword, row.get("IP地址", ""), row.get("账号", ""))
            if key in seen:
                continue
            seen.add(key)
            output.append({"查询关键字": keyword, **row, "结果": "已找到"})
    return output


def upsert(row: Dict) -> Dict:
    ip_value = _validate_ip(row.get("IP地址", ""))
    rows = read_ledger()
    normalized = {
        "IP地址": ip_value,
        "状态": (row.get("状态") or "").strip(),
        "姓名": (row.get("姓名") or "").strip(),
        "账号": (row.get("账号") or "").strip(),
        "房间号": (row.get("房间号") or "").strip(),
    }
    updated = False
    for index, existing in enumerate(rows):
        if existing.get("IP地址") == ip_value:
            rows[index] = normalized
            updated = True
            break
    if not updated:
        rows.append(normalized)
    write_ledger(rows)
    return {"row": normalized, "action": "updated" if updated else "created"}


def _read_upload_rows(path: Path) -> List[Dict]:
    suffix = path.suffix.lower()
    raw_rows: List[List[str]] = []
    if suffix == ".csv":
        raw_rows = _read_csv_rows(path)
    elif suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            raw_rows = [[str(cell.value or "").strip() for cell in row] for row in ws.iter_rows()]
        finally:
            wb.close()
    else:
        raise BenBLedgerError("仅支持 .xlsx / .csv 文件")

    if raw_rows and _looks_like_header(raw_rows[0]):
        raw_rows = raw_rows[1:]
    return [_normalize_row(row) for row in raw_rows if any(row)]


def import_rows(path: Path) -> Dict:
    imported = 0
    skipped = 0
    for row in _read_upload_rows(path):
        if not row.get("IP地址"):
            skipped += 1
            continue
        try:
            upsert(row)
            imported += 1
        except BenBLedgerError:
            skipped += 1
    return {"imported": imported, "skipped": skipped}


def _network_from_row(row: Dict):
    address = row.get("地址段", "").strip()
    mask = row.get("掩码", "").strip().lstrip("/")
    if not address:
        return None
    try:
        return ipaddress.ip_network(f"{address}/{mask}" if mask and "/" not in address else address, strict=False)
    except ValueError:
        return None


def _company_pools():
    pools = []

    for row in _read_pool_rows():
        enabled = row.get("是否启用", "").strip()
        if enabled not in {"是", "启用", "Y", "y", "yes", "YES", "1", "true", "True"}:
            continue
        network = _network_from_row(row)
        if network is not None:
            pools.append(network)

    unique = {}
    for pool in pools:
        unique[(pool.version, int(pool.network_address), pool.prefixlen)] = pool
    return sorted(unique.values(), key=lambda net: (int(net.network_address), net.prefixlen))


def _read_pool_rows() -> List[Dict]:
    rows = _read_csv_rows(POOL_FILE)
    if not rows:
        return []
    header = rows[0]
    data_rows = rows[1:] if "地址段" in "".join(header) else rows
    result = []
    for row in data_rows:
        values = list(row) + [""] * max(0, 5 - len(row))
        result.append({
            "名称": values[0].strip(),
            "地址段": values[1].strip(),
            "掩码": values[2].strip(),
            "用途": values[3].strip(),
            "是否启用": values[4].strip(),
        })
    return result


def _occupied_ips() -> Set[str]:
    occupied = set()
    for row in read_ledger():
        try:
            occupied.add(str(ipaddress.IPv4Address(row.get("IP地址", ""))))
        except ValueError:
            continue
    return occupied


def available_ips(count: int = 20) -> List[str]:
    count = max(1, min(int(count or 20), 200))
    occupied = _occupied_ips()
    available = []
    for pool in _company_pools():
        for ip in pool.hosts():
            ip_text = str(ip)
            if ip_text in occupied:
                continue
            available.append(ip_text)
            if len(available) >= count:
                return available
    return available
