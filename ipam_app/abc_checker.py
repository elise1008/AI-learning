import csv
import ipaddress
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Set, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


IP_PATTERN = re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b")


class AbcCheckError(ValueError):
    pass


def _valid_ip(value: str) -> bool:
    try:
        ipaddress.IPv4Address(value.strip())
        return True
    except ValueError:
        return False


def _extract_ips(text) -> Set[str]:
    if text is None:
        return set()
    return {ip for ip in IP_PATTERN.findall(str(text)) if _valid_ip(ip)}


def extract_ips_from_file(path: Path) -> Set[str]:
    suffix = path.suffix.lower()
    ips: Set[str] = set()

    if suffix == ".csv":
        for encoding in ("utf-8-sig", "gbk"):
            try:
                with open(path, "r", encoding=encoding, newline="") as f:
                    reader = csv.reader(f)
                    for row in reader:
                        for cell in row:
                            ips.update(_extract_ips(cell))
                return ips
            except UnicodeDecodeError:
                continue
        raise AbcCheckError("CSV 文件编码无法识别")

    if suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        ips.update(_extract_ips(cell.value))
        finally:
            wb.close()
        return ips

    raise AbcCheckError("仅支持 .xlsx / .csv 文件")


def _sort_ip(ip: str) -> int:
    return int(ipaddress.IPv4Address(ip))


def build_results(a_path: Path, b_path: Path, c_path: Path) -> Tuple[List[Dict], Set[str], Set[str], Set[str]]:
    a_set = extract_ips_from_file(a_path)
    b_set = extract_ips_from_file(b_path)
    c_set = extract_ips_from_file(c_path)

    all_ips = sorted(a_set | b_set | c_set, key=_sort_ip)
    results = []
    for ip in all_ips:
        compliant = ip in a_set and ip in b_set and ip in c_set
        results.append({
            "IP地址": ip,
            "桌管": "已安装" if ip in a_set else "未安装",
            "V10": "已安装" if ip in b_set else "未安装",
            "合规性": "已安装" if ip in c_set else "未安装",
            "合规状态": "合规" if compliant else "不合规",
        })
    return results, a_set, b_set, c_set


def create_abc_check_workbook(a_path: Path, b_path: Path, c_path: Path, output_dir: Path) -> Dict:
    results, a_set, b_set, c_set = build_results(a_path, b_path, c_path)
    if not results:
        raise AbcCheckError("三个文件中均未检测到有效 IPv4 地址")

    wb = Workbook()
    ws = wb.active
    ws.title = "桌管_V10_合规性安装汇总"

    headers = ["IP地址", "桌管", "V10", "合规性", "合规状态"]
    ws.append(headers)
    for row in results:
        ws.append([row[header] for header in headers])

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True)
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for cell in ws["E"][1:]:
        cell.fill = pass_fill if cell.value == "合规" else fail_fill

    ws.freeze_panes = "A2"
    for index, width in enumerate([18, 12, 12, 12, 14], 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ref = f"A1:E{ws.max_row}"
    table = Table(displayName="AbcComplianceSummary", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"桌管_V10_合规性安装汇总_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(output_path)

    compliant = sum(1 for row in results if row["合规状态"] == "合规")
    return {
        "output_path": output_path,
        "total": len(results),
        "compliant": compliant,
        "non_compliant": len(results) - compliant,
        "a_count": len(a_set),
        "b_count": len(b_set),
        "c_count": len(c_set),
    }
