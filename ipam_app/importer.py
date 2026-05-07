"""
导入模块
支持 Excel(xlsx) 和 CSV 文件导入
自动识别字段，批量导入，重复提示，错误日志
"""
import csv
from pathlib import Path
from typing import Dict, List
import config
import data_manager


try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def import_file(file_path: str, target_biz: str) -> Dict:
    """
    导入 Excel 或 CSV 文件到指定业务表
    返回: {"success": True/False, "imported": 0, "skipped": 0, "errors": [...], "message": "..."}
    """
    path = Path(file_path)
    if not path.exists():
        return {"success": False, "imported": 0, "skipped": 0, "errors": [f"文件不存在: {file_path}"], "message": "文件不存在"}

    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _read_csv_file(path)
    elif suffix in (".xlsx", ".xls"):
        if not HAS_OPENPYXL:
            return {"success": False, "imported": 0, "skipped": 0, "errors": ["缺少 openpyxl 库"], "message": "请安装 openpyxl: pip install openpyxl"}
        rows = _read_excel_file(path)
    else:
        return {"success": False, "imported": 0, "skipped": 0, "errors": ["不支持的格式"], "message": "仅支持 xlsx/csv 格式"}

    if not rows:
        return {"success": False, "imported": 0, "skipped": 0, "errors": ["文件为空"], "message": "文件中没有数据"}

    columns = config.BUSINESS_COLUMNS.get(target_biz, list(rows[0].keys()))

    if target_biz == "master_pool":
        dest_file = config.MASTER_POOL_FILE
    elif target_biz == "router_loopback":
        dest_file = config.ROUTER_LOOPBACK_FILE
    elif target_biz == "router_link":
        dest_file = config.ROUTER_LINK_FILE
    else:
        dest_file = config.BUSINESS_FILES.get(target_biz)

    if not dest_file:
        return {"success": False, "imported": 0, "skipped": 0, "errors": [f"目标业务不存在: {target_biz}"], "message": "目标业务不存在"}

    existing = data_manager.read_csv(dest_file)
    existing_devices = {r.get("设备名称", "").strip(): r for r in existing}

    imported = 0
    skipped = 0
    errors = []

    for i, row in enumerate(rows, start=2):
        normalized = _normalize_row(row, columns)
        device_name = normalized.get("设备名称", "").strip()

        if not device_name:
            errors.append(f"第{i}行: 设备名称为空，已跳过")
            skipped += 1
            continue

        if device_name in existing_devices:
            errors.append(f"第{i}行: 设备 {device_name} 已存在，已跳过")
            skipped += 1
            continue

        existing.append(normalized)
        existing_devices[device_name] = normalized
        imported += 1

    data_manager.write_csv(dest_file, existing, columns)

    return {
        "success": True,
        "imported": imported,
        "skipped": skipped,
        "errors": errors,
        "message": f"导入完成：成功 {imported} 条，跳过 {skipped} 条"
    }


def _read_csv_file(file_path: Path) -> List[Dict]:
    """读取CSV"""
    try:
        with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            return [row for row in reader]
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="gbk", newline="") as f:
            reader = csv.DictReader(f)
            return [row for row in reader]
    except Exception as e:
        return []


def _read_excel_file(file_path: Path) -> List[Dict]:
    """读取Excel"""
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.worksheets[0]
        rows = []
        headers = [str(cell.value).strip() if cell.value else f"_col{i}" for i, cell in enumerate(sheet[1], start=1)]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for j, cell in enumerate(row):
                if j < len(headers):
                    row_dict[headers[j]] = str(cell).strip() if cell is not None else ""
            rows.append(row_dict)
        return rows
    except Exception as e:
        return []


def read_csv_from_fileobj(fileobj) -> List[Dict]:
    """从 Flask file 对象读取 CSV"""
    try:
        content = fileobj.stream.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        fileobj.stream.seek(0)
        content = fileobj.stream.read().decode("gbk")
    except Exception:
        return []

    try:
        lines = content.splitlines()
        if not lines:
            return []
        reader = csv.DictReader(lines)
        return [row for row in reader]
    except Exception:
        return []


def read_excel_from_fileobj(fileobj) -> List[Dict]:
    """从 Flask file 对象读取 Excel"""
    if not HAS_OPENPYXL:
        return []
    try:
        wb = openpyxl.load_workbook(fileobj.stream, data_only=True)
        sheet = wb.worksheets[0]
        rows = []
        headers = [str(cell.value).strip() if cell.value else f"_col{i}" for i, cell in enumerate(sheet[1], start=1)]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for j, cell in enumerate(row):
                if j < len(headers):
                    row_dict[headers[j]] = str(cell).strip() if cell is not None else ""
            rows.append(row_dict)
        return rows
    except Exception:
        return []


def _normalize_row(row: Dict, columns: List[str]) -> Dict:
    """规范化行数据，匹配目标列"""
    normalized = {col: "" for col in columns}
    for key, value in row.items():
        key_stripped = key.strip()
        value_stripped = str(value).strip() if value else ""
        if key_stripped in columns:
            normalized[key_stripped] = value_stripped
        else:
            for col in columns:
                if col in key_stripped or key_stripped in col:
                    normalized[col] = value_stripped
                    break
    return normalized
