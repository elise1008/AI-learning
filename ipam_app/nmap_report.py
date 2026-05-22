import re
import ipaddress
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

import config
import data_manager


TARGET_PORTS = [21, 22, 135, 137, 138, 139, 445, 3389]
IPAM_BUSINESS_SCOPE = "xxnw"
SERVICE_NAMES = {
    21: "FTP",
    22: "SSH",
    135: "MSRPC",
    137: "NetBIOS Name",
    138: "NetBIOS Datagram",
    139: "NetBIOS Session",
    445: "SMB",
    3389: "RDP",
}


class NmapReportError(ValueError):
    pass


def _sort_ip(ip: str) -> Tuple[int, ...]:
    try:
        return tuple(int(part) for part in ip.split("."))
    except ValueError:
        return (999, 999, 999, 999)


def _empty_host(ip: str) -> Dict:
    return {
        "ip": ip,
        "ports": {port: "未出现" for port in TARGET_PORTS},
    }


def parse_nmap_text(text: str) -> List[Dict]:
    if not text.strip():
        raise NmapReportError("上传的 nmap 文件为空")

    hosts: Dict[str, Dict] = {}
    current_ip = None

    for raw_line in text.splitlines():
        line = raw_line.strip()
        report_match = re.match(r"^Nmap scan report for\s+(.+)$", line, re.I)
        if report_match:
            target = report_match.group(1).strip()
            ip_match = re.search(r"(\d{1,3}(?:\.\d{1,3}){3})", target)
            current_ip = ip_match.group(1) if ip_match else target
            hosts.setdefault(current_ip, _empty_host(current_ip))
            continue

        if not current_ip:
            continue

        port_match = re.match(r"^(\d+)/(tcp|udp)\s+(\S+)\s*(.*)$", line, re.I)
        if not port_match:
            continue

        port = int(port_match.group(1))
        if port in TARGET_PORTS:
            hosts[current_ip]["ports"][port] = port_match.group(3).lower()

    if not hosts:
        raise NmapReportError("未识别到 nmap 扫描结果")

    return sorted(hosts.values(), key=lambda host: _sort_ip(host["ip"]))


def _open_ports(host: Dict) -> List[int]:
    return [port for port in TARGET_PORTS if host["ports"].get(port) == "open"]


def _filtered_ports(host: Dict) -> List[int]:
    return [port for port in TARGET_PORTS if host["ports"].get(port) == "filtered"]


def _port_label(port: int) -> str:
    return f"{port}/{SERVICE_NAMES[port]}"


def _parse_network(address: str, mask: str):
    address = (address or "").strip()
    mask = (mask or "").strip().lstrip("/")
    if not address:
        return None

    try:
        if "/" in address:
            return ipaddress.ip_network(address, strict=False)
        if mask:
            return ipaddress.ip_network(f"{address}/{mask}", strict=False)
        return ipaddress.ip_network(f"{address}/32", strict=False)
    except ValueError:
        return None


def _build_ipam_map(ips: List[str]) -> Dict[str, Dict]:
    """Only match Nmap IPs against the xxnw business segment."""
    if not ips:
        return {}

    file_path = config.BUSINESS_FILES.get(IPAM_BUSINESS_SCOPE)
    rows = data_manager.read_csv(file_path) if file_path else []
    networks = []
    for row in rows:
        network = _parse_network(row.get("地址段", ""), row.get("掩码", ""))
        if network is not None:
            networks.append((network, row))

    result = {}
    for ip in ips:
        try:
            address = ipaddress.ip_address(ip)
        except ValueError:
            result[ip] = {
                "查询IP": ip,
                "归属设备": "无效IP",
                "所属业务": IPAM_BUSINESS_SCOPE,
                "地址段": "",
                "掩码": "",
                "备注": "格式错误",
            }
            continue

        matched = None
        for network, row in networks:
            if address in network:
                matched = row
                break

        if matched:
            result[ip] = {
                "查询IP": ip,
                "归属设备": matched.get("设备名称", ""),
                "所属业务": IPAM_BUSINESS_SCOPE,
                "地址段": matched.get("地址段", ""),
                "掩码": matched.get("掩码", ""),
                "备注": "已分配",
            }
        else:
            result[ip] = {
                "查询IP": ip,
                "归属设备": "无归属",
                "所属业务": IPAM_BUSINESS_SCOPE,
                "地址段": "",
                "掩码": "",
                "备注": "未分配",
            }

    return result


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        width = min(max(max_len + 2, 12), 42)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def _add_table(ws, table_name: str) -> None:
    if ws.max_row < 1 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def create_nmap_risk_workbook(text: str, output_dir: Path) -> Dict:
    hosts = parse_nmap_text(text)
    open_hosts = [host for host in hosts if _open_ports(host)]
    ipam_map = _build_ipam_map([host["ip"] for host in open_hosts])

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "开放高危端口汇总"
    status_ws = wb.create_sheet("全部目标端口状态")

    summary_ws.append([
        "IP",
        "存在的开放高危端口",
        "开放端口数量",
        "被过滤的目标端口",
        "归属设备",
        "所属业务",
        "IPAM地址段",
        "IPAM掩码",
        "IPAM备注",
    ])
    for host in open_hosts:
        ip = host["ip"]
        ipam = ipam_map.get(ip, {})
        summary_ws.append([
            ip,
            ", ".join(_port_label(port) for port in _open_ports(host)),
            len(_open_ports(host)),
            ", ".join(_port_label(port) for port in _filtered_ports(host)),
            ipam.get("归属设备", ""),
            ipam.get("所属业务", ""),
            ipam.get("地址段", ""),
            ipam.get("掩码", ""),
            ipam.get("备注", ""),
        ])

    status_ws.append([
        "IP",
        "开放高危端口",
        *[_port_label(port) for port in TARGET_PORTS],
    ])
    for host in hosts:
        status_ws.append([
            host["ip"],
            ", ".join(str(port) for port in _open_ports(host)),
            *[host["ports"].get(port, "未出现") for port in TARGET_PORTS],
        ])

    _style_sheet(summary_ws)
    _style_sheet(status_ws)
    _add_table(summary_ws, "OpenRiskPorts")
    _add_table(status_ws, "TargetPortStatus")

    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"nmap_高危端口归属_{timestamp}.xlsx"
    wb.save(output_path)

    return {
        "output_path": output_path,
        "scanned_hosts": len(hosts),
        "open_risk_hosts": len(open_hosts),
    }
